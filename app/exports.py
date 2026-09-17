"""CSV / Excel / PDF export helpers."""
import csv
import io
from datetime import date

from flask import Response, current_app


def to_csv(filename, headers, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return Response(
        buf.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
    )


def to_excel(filename, headers, rows, title=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (title or filename)[:31]

    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append(list(row))

    for i, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in rows:
            if i <= len(row):
                longest = max(longest, len(str(row[i - 1] if row[i - 1] is not None else "")))
        ws.column_dimensions[get_column_letter(i)].width = min(50, max(10, longest + 2))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
    )


def to_pdf(filename, title, headers, rows, subtitle=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    navy = colors.HexColor("#1F3A5F")

    title_style = ParagraphStyle("t", parent=styles["Heading1"], fontSize=15, textColor=navy, spaceAfter=2)
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=9,
                               textColor=colors.HexColor("#5B6B7C"), spaceAfter=8)
    cell_style = ParagraphStyle("c", parent=styles["Normal"], fontSize=7.5, leading=9.5)
    head_style = ParagraphStyle("h", parent=styles["Normal"], fontSize=7.5, leading=9.5,
                                textColor=colors.white, fontName="Helvetica-Bold")

    story = [Paragraph(title, title_style)]
    meta = f"{current_app.config['ORG_NAME']} · generated {date.today().strftime('%d %b %Y')}"
    story.append(Paragraph(subtitle + " · " + meta if subtitle else meta, sub_style))
    story.append(Spacer(1, 4))

    data = [[Paragraph(str(h), head_style) for h in headers]]
    for row in rows:
        data.append([Paragraph("" if v is None else str(v), cell_style) for v in row])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D2DB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F6")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    doc.build(story)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}.pdf"},
    )


def export_response(fmt, filename, title, headers, rows, subtitle=None):
    """Dispatch to the requested export format."""
    if fmt == "xlsx":
        return to_excel(filename, headers, rows, title=title)
    if fmt == "pdf":
        return to_pdf(filename, title, headers, rows, subtitle=subtitle)
    return to_csv(filename, headers, rows)
